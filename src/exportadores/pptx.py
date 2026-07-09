import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

# ─────────────────────────────────────────────────────
# LOGIC DE DIRETÓRIOS DINÂMICOS
# ─────────────────────────────────────────────────────
DIRETORIO_EXPORTADOR = os.path.dirname(os.path.abspath(__file__))
PASTA_SRC = os.path.abspath(os.path.join(DIRETORIO_EXPORTADOR, ".."))
PASTA_RAIZ_PROJETO = os.path.abspath(os.path.join(PASTA_SRC, ".."))
PASTA_ASSETS = os.path.join(PASTA_RAIZ_PROJETO, "assets")

print(f" > Buscando assets em: {PASTA_ASSETS}")


class ExportadorPPTX:
    @staticmethod
    def gerar_layout(pedido, nome_arquivo):
        prs = Presentation()
        prs.slide_width = Cm(40)
        prs.slide_height = Cm(22.5)  # Proporção Widescreen 16:9

        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # --- 1. CABEÇALHO (Logo, Textos e Linha) ---
        caminho_logo = os.path.join(PASTA_ASSETS, "logo.png")
        if os.path.exists(caminho_logo):
            slide.shapes.add_picture(caminho_logo, Cm(1), Cm(0.5), height=Cm(2.5))

        # Texto descritivo abaixo da logo
        tx_desc = slide.shapes.add_textbox(Cm(0.8), Cm(3), Cm(15), Cm(1.5))
        tf_desc = tx_desc.text_frame
        tf_desc.word_wrap = True

        p1 = tf_desc.paragraphs[0]
        p1.text = "Linha Definitiva Solução em Travessas para Self-Service"
        p1.font.bold = True
        p1.font.size = Pt(11)

        p2 = tf_desc.add_paragraph()
        p2.text = "@Bonanza_Travessasebalcoes 31-99084-4400"
        p2.font.size = Pt(11)

        # Linha reta cobrindo o texto
        linha = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Cm(1), Cm(4.5), Cm(12.5), Cm(4.5)
        )
        linha.line.color.rgb = RGBColor(0, 0, 0)
        linha.line.width = Pt(1.5)

        # Info do Cliente
        tx_cliente = slide.shapes.add_textbox(Cm(1), Cm(4.8), Cm(15), Cm(1))
        p_cliente = tx_cliente.text_frame.paragraphs[0]
        p_cliente.text = f"Cliente: {pedido.nome_cliente}"
        p_cliente.font.bold = True
        p_cliente.font.size = Pt(12)

        # Título Centralizado
        tx_titulo = slide.shapes.add_textbox(Cm(10), Cm(1.5), Cm(20), Cm(2))
        tf_titulo = tx_titulo.text_frame
        tf_titulo.text = "Opção 1\nSelf-Service Quente"
        tf_titulo.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf_titulo.paragraphs[0].font.bold = True
        tf_titulo.paragraphs[0].font.size = Pt(20)
        tf_titulo.paragraphs[0].font.underline = True
        if len(tf_titulo.paragraphs) > 1:
            tf_titulo.paragraphs[1].alignment = PP_ALIGN.CENTER
            tf_titulo.paragraphs[1].font.bold = True
            tf_titulo.paragraphs[1].font.size = Pt(18)
            tf_titulo.paragraphs[1].font.underline = True
            tf_titulo.paragraphs[1].font.italic = True

        # --- 2. DEFINIÇÃO DA ESCALA (1:10) ---
        ESCALA = 0.1

        for modulo in pedido.modulos:
            largura_real = Cm(modulo.engine.L * ESCALA)
            altura_real = Cm(modulo.engine.P * ESCALA)

            # Centralização dinâmica baseada na escala do slide inteiro
            offset_x = (prs.slide_width - largura_real) / 2
            offset_y = Cm(7)

            # --- 3. DESENHO DO BALCÃO ---
            caminho_fundo = os.path.join(PASTA_ASSETS, "fundo_perfurado.png")
            if os.path.exists(caminho_fundo):
                slide.shapes.add_picture(
                    caminho_fundo,
                    offset_x,
                    offset_y,
                    largura_real,
                    altura_real,
                )
            else:
                fundo = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE, offset_x, offset_y, largura_real, altura_real
                )
                fundo.fill.solid()
                fundo.fill.fore_color.rgb = RGBColor(210, 210, 210)

            # --- 4. ALOCAÇÃO DAS PANELAS ---
            for item in modulo.engine.itens:
                x_pos = offset_x + Cm(item["x"] * ESCALA)
                y_pos = offset_y + Cm(item["y"] * ESCALA)
                w_dim = Cm(item["w"] * ESCALA)
                h_dim = Cm(item["h"] * ESCALA)

                is_circulo = item.get("formato") == "circulo"

                if is_circulo:
                    w_original = w_dim
                    h_original = h_dim
                    diametro_base = min(item["w"], item["h"]) * ESCALA
                    h_dim = Cm(diametro_base * 0.95)
                    w_dim = Cm(diametro_base * 0.95 * 1.08)
                    x_pos = x_pos + (w_original - w_dim) / 2
                    y_pos = y_pos + (h_original - h_dim) / 2

                # PASSO 1: Criar o Grupo no Slide
                grupo = slide.shapes.add_group_shape()

                caminho_panela_redonda = os.path.join(
                    PASTA_ASSETS, "panela_redonda.png"
                )
                caminho_panela_retangular = os.path.join(
                    PASTA_ASSETS, "panela_retangular.png"
                )
                img_path = (
                    caminho_panela_redonda if is_circulo else caminho_panela_retangular
                )

                # PASSO 2: Adicionar a Travessa DENTRO do Grupo
                shape_travessa = None
                if os.path.exists(img_path):
                    shape_travessa = grupo.shapes.add_picture(
                        img_path, x_pos, y_pos, w_dim, h_dim
                    )
                    if is_circulo:
                        shape_travessa.rotation = 33.8
                else:
                    forma_fallback = (
                        MSO_SHAPE.OVAL if is_circulo else MSO_SHAPE.ROUNDED_RECTANGLE
                    )
                    shape_travessa = grupo.shapes.add_shape(
                        forma_fallback, x_pos, y_pos, w_dim, h_dim
                    )
                    shape_travessa.fill.solid()
                    shape_travessa.fill.fore_color.rgb = RGBColor(255, 255, 255)
                    if is_circulo:
                        shape_travessa.rotation = 33.8

                # PASSO 3: Adicionar o Texto DENTRO do Grupo
                tx_item = grupo.shapes.add_textbox(x_pos, y_pos, w_dim, h_dim)
                tf_item = tx_item.text_frame
                tf_item.word_wrap = False
                tf_item.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf_item.margin_top = tf_item.margin_bottom = tf_item.margin_left = (
                    tf_item.margin_right
                ) = Cm(0)

                p_item = tf_item.paragraphs[0]
                p_item.text = (
                    f"{item['w']}" if is_circulo else f"{item['w']}x{item['h']}"
                )
                p_item.alignment = PP_ALIGN.CENTER
                p_item.font.bold = True
                p_item.font.size = Pt(13)
                p_item.font.color.rgb = RGBColor(0, 0, 0)

                grupo.name = f"Travessa_{item['nome']}"

            # --- 5. MARCADORES DE DIMENSÃO ---

            # 5.1. Chave Inferior (Largura)
            target_x = offset_x
            target_y = offset_y + altura_real + Cm(0.2)
            target_w = largura_real
            target_h = Cm(0.8)

            cx = target_x + target_w / 2
            cy = target_y + target_h / 2

            chave_inf = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_BRACE,
                cx - target_h / 2,
                cy - target_w / 2,
                target_h,
                target_w,
            )
            chave_inf.rotation = 90
            chave_inf.fill.background()
            chave_inf.line.color.rgb = RGBColor(0, 0, 0)
            chave_inf.line.width = Pt(1.5)

            tx_largura = slide.shapes.add_textbox(
                offset_x, target_y + Cm(0.6), largura_real, Cm(1.5)
            )
            p_larg = tx_largura.text_frame.paragraphs[0]
            p_larg.text = f"{modulo.engine.L}cm"
            p_larg.alignment = PP_ALIGN.CENTER
            p_larg.font.bold = True
            p_larg.font.size = Pt(16)

            # 5.2. Chave Lateral Direita (Profundidade)
            chave_dir = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_BRACE,
                offset_x + largura_real + Cm(0.2),
                offset_y,
                Cm(0.8),
                altura_real,
            )
            chave_dir.fill.background()
            chave_dir.line.color.rgb = RGBColor(0, 0, 0)
            chave_dir.line.width = Pt(1.5)

            tx_prof = slide.shapes.add_textbox(
                offset_x + largura_real + Cm(1.0), offset_y, Cm(3), altura_real
            )
            tf_prof = tx_prof.text_frame
            tf_prof.word_wrap = False
            tf_prof.vertical_anchor = MSO_ANCHOR.MIDDLE

            tf_prof.margin_top = Cm(0)
            tf_prof.margin_bottom = Cm(0)
            tf_prof.margin_left = Cm(0.1)
            tf_prof.margin_right = Cm(0)

            p_prof = tx_prof.text_frame.paragraphs[0]
            p_prof.text = f"{modulo.engine.P}cm"
            p_prof.alignment = PP_ALIGN.LEFT
            p_prof.font.bold = True
            p_prof.font.size = Pt(16)

        prs.save(nome_arquivo)
        print(f" > Layout PPTX centralizado e exportado com sucesso: '{nome_arquivo}'")
